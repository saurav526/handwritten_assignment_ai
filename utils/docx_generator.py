from docx import Document

def create_docx(title: str, content: str, output_path: str):
    doc = Document()
    doc.add_heading(title, level=1)

    for line in content.splitlines():
        line = line.strip()
        if not line:
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        else:
            doc.add_paragraph(line)

    doc.save(output_path)
